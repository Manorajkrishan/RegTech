"""Download official revised 2026 UK factors into a REVIEW-ONLY raw catalogue.

This does not silently replace the active, vetted emission-factor mapping.
An accountable reviewer must inspect units, scopes, factor variants and source
rows before publishing a production factor mapping.
"""
import hashlib
import io
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

URL = ("https://assets.publishing.service.gov.uk/media/"
       "6a6c9748862aaf18d9c62ac9/"
       "ghg-conversion-factors-2026-flat-format-revised.xlsx")
GOV_PAGE = "https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2026"
DATA_DIR = Path(__file__).resolve().parent.parent / "data"
MAX_BYTES = 10_000_000


def safe_value(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def export_catalogue(contents: bytes, output: Path):
    book = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    sheets = {}
    for sheet in book:
        # Preserve worksheet rows and original positions rather than guessing
        # which of the several CO2/CH4/N2O/WTT columns a metric requires.
        rows = []
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            row = [safe_value(value) for value in values]
            if any(value is not None for value in row):
                rows.append({"source_row": index, "cells": row})
        sheets[sheet.title] = rows
    payload = {
        "year": 2026, "status": "unreviewed_import_not_approved_for_calculation",
        "source_page": GOV_PAGE, "file_url": URL,
        "revision": "31 July 2026",
        "retrieved_utc": datetime.now(timezone.utc).isoformat(),
        "sha256": hashlib.sha256(contents).hexdigest(),
        "sheets": sheets,
    }
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return {name: len(rows) for name, rows in sheets.items()}


if __name__ == "__main__":
    DATA_DIR.mkdir(exist_ok=True)
    request = urllib.request.Request(URL, headers={"User-Agent": "RegTech-data-sync/0.2"})
    with urllib.request.urlopen(request, timeout=45) as response:
        contents = response.read(MAX_BYTES + 1)
    if len(contents) > MAX_BYTES:
        raise SystemExit("Unexpectedly large workbook")
    output = DATA_DIR / "defra_2026_raw_catalogue.json"
    stats = export_catalogue(contents, output)
    print("Source", GOV_PAGE)
    print("SHA-256:", hashlib.sha256(contents).hexdigest())
    print("Raw factors exported to", output)
    print("Rows per sheet:", stats)
    print("REVIEW REQUIRED: do not use this unreviewed catalogue directly as emission factors.")
